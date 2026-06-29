"""Batch mentor extraction through the local OpenAI-compatible endpoint."""

from __future__ import annotations

import argparse
import json
import os
import re
import time
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Literal, Mapping

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from mentor_agent.excel_io import ExcelLoadReport, load_mentor_inputs_with_report
from mentor_agent.schemas import MentorInput, MentorResult
from mentor_agent.simple_schemas import SimpleMentorBatchResult, SimpleMentorResult
from mentor_agent.tag_taxonomy import load_tag_taxonomy


DEFAULT_BASE_URL = "http://127.0.0.1:9000"
DEFAULT_OUTPUT_DIR = Path("outputs")
SUCCESS_CHECKPOINT_RELATIVE = Path("checkpoints") / "success_rows.jsonl"
FINAL_RESULTS_NAME = "mentor_results.jsonl"
FAILED_ROWS_NAME = "failed_rows.jsonl"
REVIEW_EXCEL_NAME = "mentor_review.xlsx"
EXTRACTION_MODES = ("simple", "full")

ExtractionMode = Literal["simple", "full"]
ExtractionResult = SimpleMentorResult | MentorResult

REVIEW_SHEETS = (
    "导师总览",
    "行业标签",
    "任职经历",
    "技能",
    "证书资质奖项",
    "教育背景",
    "辅导人群",
    "职业亮点",
    "质量问题",
    "处理失败",
)
SIMPLE_REVIEW_SHEETS = ("导师总览", "关键词明细", "处理失败")
TAGGED_SIMPLE_REVIEW_SHEETS = (
    "导师总览",
    "关键词明细",
    "标准行业标签",
    "标准职位标签",
    "公司标签",
    "人工审核项",
    "处理失败",
)

_SENSITIVE_PATTERN = re.compile(
    r"(?i)(access\s*key|accesskey|api\s*key|authorization|\.env|token)"
    r"(\s*[:=]\s*|\s+)[^\s,;]+"
)


@dataclass(frozen=True)
class BatchConfig:
    input_path: Path
    output_dir: Path = DEFAULT_OUTPUT_DIR
    base_url: str = DEFAULT_BASE_URL
    limit: int | None = None
    offset: int = 0
    resume: bool = False
    force: bool = False
    dry_run: bool = False
    timeout: float = 120.0
    max_retries: int = 2
    retry_sleep: float = 1.0
    no_excel: bool = False
    verbose: bool = False
    mode: ExtractionMode = "simple"
    batch_size: int = 10
    standard_tags_enabled: bool = False
    taxonomy_path: Path = Path("configs/职位类型_2.txt")


@dataclass(frozen=True)
class BatchSummary:
    source_file: str
    valid_mentor_count: int
    selected_count: int
    success_count: int
    failed_count: int
    skipped_resume_count: int
    final_result_count: int
    dry_run: bool

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class BatchItemError(RuntimeError):
    def __init__(
        self,
        error_type: str,
        message: str,
        *,
        attempt_count: int,
        retriable: bool,
        stage: str = "http_extract",
    ) -> None:
        super().__init__(message)
        self.error_type = error_type
        self.safe_message = sanitize_error_message(message)
        self.attempt_count = attempt_count
        self.retriable = retriable
        self.stage = stage


def sanitize_error_message(message: str, max_length: int = 200) -> str:
    sanitized = _SENSITIVE_PATTERN.sub(r"\1=<redacted>", message)
    sanitized = " ".join(sanitized.split())
    if len(sanitized) > max_length:
        return sanitized[: max_length - 3] + "..."
    return sanitized


def _enum_value(value: Any) -> Any:
    return getattr(value, "value", value)


def _non_negative_int(value: str) -> int:
    parsed = int(value)
    if parsed < 0:
        raise argparse.ArgumentTypeError("must be >= 0")
    return parsed


def _positive_float(value: str) -> float:
    parsed = float(value)
    if parsed <= 0:
        raise argparse.ArgumentTypeError("must be > 0")
    return parsed


def _non_negative_float(value: str) -> float:
    parsed = float(value)
    if parsed < 0:
        raise argparse.ArgumentTypeError("must be >= 0")
    return parsed


def _positive_int(value: str) -> int:
    parsed = int(value)
    if parsed <= 0:
        raise argparse.ArgumentTypeError("must be > 0")
    return parsed


def _boolean_env(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    normalized = value.strip().lower()
    if normalized == "true":
        return True
    if normalized == "false":
        return False
    raise ValueError(f"{name} must be 'true' or 'false'")


def _endpoint(base_url: str) -> str:
    return base_url.rstrip("/") + "/openai/v1/chat/completions"


def _response_error(
    error_type: str,
    message: str,
    attempt_count: int,
    *,
    retriable: bool = True,
) -> BatchItemError:
    return BatchItemError(
        error_type,
        message,
        attempt_count=attempt_count,
        retriable=retriable,
    )


def _result_model_for_mode(mode: ExtractionMode) -> type[ExtractionResult]:
    if mode == "simple":
        return SimpleMentorResult
    if mode == "full":
        return MentorResult
    raise ValueError("mode must be 'simple' or 'full'")


def _validate_standard_tag_metadata(
    result: SimpleMentorResult,
    *,
    expected_enabled: bool,
    expected_taxonomy_hash: str | None,
    attempt_count: int,
) -> None:
    if result.standard_tags_enabled != expected_enabled:
        raise _response_error(
            "standard_tags_mode_mismatch",
            "service standard-tags mode did not match batch configuration",
            attempt_count,
            retriable=False,
        )
    if expected_enabled and result.taxonomy_hash != expected_taxonomy_hash:
        raise _response_error(
            "taxonomy_hash_mismatch",
            "service taxonomy hash did not match batch taxonomy",
            attempt_count,
            retriable=False,
        )


def _parse_http_response(
    response: Any,
    mentor_input: MentorInput,
    attempt_count: int,
    *,
    mode: ExtractionMode,
    standard_tags_enabled: bool = False,
    taxonomy_hash: str | None = None,
) -> ExtractionResult:
    status_code = int(response.status_code)
    if not 200 <= status_code < 300:
        retriable = status_code == 429 or status_code >= 500
        raise _response_error(
            "http_status_error",
            f"local service returned HTTP {status_code}",
            attempt_count,
            retriable=retriable,
        )

    try:
        payload = response.json()
    except (ValueError, json.JSONDecodeError):
        raise _response_error(
            "invalid_response_shape",
            "local service response was not valid JSON",
            attempt_count,
        ) from None

    if not isinstance(payload, Mapping):
        raise _response_error(
            "invalid_response_shape",
            "local service JSON response was not an object",
            attempt_count,
        )
    if "error" in payload:
        raise _response_error(
            "invalid_response_shape",
            "local service returned an error object",
            attempt_count,
        )

    try:
        content = payload["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError):
        raise _response_error(
            "missing_message_content",
            "response did not contain choices[0].message.content",
            attempt_count,
        ) from None

    if not isinstance(content, str) or not content.strip():
        raise _response_error(
            "missing_message_content",
            "choices[0].message.content was empty or not text",
            attempt_count,
        )

    try:
        result_payload = json.loads(content)
    except json.JSONDecodeError:
        raise _response_error(
            "invalid_mentor_result_json",
            "message content was not valid mentor result JSON",
            attempt_count,
        ) from None

    result_model = _result_model_for_mode(mode)
    try:
        result = result_model.model_validate(result_payload)
    except ValidationError as exc:
        raise _response_error(
            "mentor_result_validation_failed",
            f"{result_model.__name__} failed schema validation ({exc.error_count()} error(s))",
            attempt_count,
        ) from None

    if (
        result.mentor_id != mentor_input.mentor_id
        or result.record_hash != mentor_input.record_hash
    ):
        raise _response_error(
            "mentor_result_validation_failed",
            "MentorResult identity did not match the requested MentorInput",
            attempt_count,
        )
    if isinstance(result, SimpleMentorResult):
        _validate_standard_tag_metadata(
            result,
            expected_enabled=standard_tags_enabled,
            expected_taxonomy_hash=taxonomy_hash,
            attempt_count=attempt_count,
        )
    return result


def _parse_http_batch_response(
    response: Any,
    mentor_inputs: list[MentorInput],
    attempt_count: int,
    *,
    standard_tags_enabled: bool = False,
    taxonomy_hash: str | None = None,
) -> list[SimpleMentorResult]:
    status_code = int(response.status_code)
    if not 200 <= status_code < 300:
        retriable = status_code == 429 or status_code >= 500
        raise _response_error(
            "http_status_error",
            f"local service returned HTTP {status_code}",
            attempt_count,
            retriable=retriable,
        )

    try:
        payload = response.json()
    except (ValueError, json.JSONDecodeError):
        raise _response_error(
            "invalid_response_shape",
            "local service response was not valid JSON",
            attempt_count,
        ) from None

    if not isinstance(payload, Mapping):
        raise _response_error(
            "invalid_response_shape",
            "local service JSON response was not an object",
            attempt_count,
        )
    if "error" in payload:
        raise _response_error(
            "invalid_response_shape",
            "local service returned an error object",
            attempt_count,
        )

    try:
        content = payload["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError):
        raise _response_error(
            "missing_message_content",
            "response did not contain choices[0].message.content",
            attempt_count,
        ) from None

    if not isinstance(content, str) or not content.strip():
        raise _response_error(
            "missing_message_content",
            "choices[0].message.content was empty or not text",
            attempt_count,
        )

    try:
        result_payload = json.loads(content)
    except json.JSONDecodeError:
        raise _response_error(
            "invalid_mentor_result_json",
            "message content was not valid simple batch result JSON",
            attempt_count,
        ) from None

    try:
        batch_result = SimpleMentorBatchResult.model_validate(result_payload)
    except ValidationError as exc:
        raise _response_error(
            "mentor_result_validation_failed",
            f"SimpleMentorBatchResult failed schema validation ({exc.error_count()} error(s))",
            attempt_count,
        ) from None

    expected = {
        (mentor_input.mentor_id, mentor_input.record_hash)
        for mentor_input in mentor_inputs
    }
    actual = {
        (result.mentor_id, result.record_hash)
        for result in batch_result.results
    }
    if actual != expected or len(batch_result.results) != len(mentor_inputs):
        raise _response_error(
            "mentor_result_validation_failed",
            "SimpleMentorBatchResult identities did not match requested MentorInputs",
            attempt_count,
        )
    for result in batch_result.results:
        _validate_standard_tag_metadata(
            result,
            expected_enabled=standard_tags_enabled,
            expected_taxonomy_hash=taxonomy_hash,
            attempt_count=attempt_count,
        )
    return batch_result.results


def request_mentor_result(
    mentor_input: MentorInput,
    *,
    base_url: str,
    timeout: float,
    max_retries: int,
    retry_sleep: float,
    session: Any,
    mode: ExtractionMode,
    standard_tags_enabled: bool = False,
    taxonomy_hash: str | None = None,
    sleep_fn: Callable[[float], None] = time.sleep,
) -> ExtractionResult:
    """Call one mentor endpoint with bounded retries and classified failures."""

    request_body = {
        "messages": [
            {
                "role": "user",
                "content": mentor_input.model_dump_json(by_alias=True),
            }
        ],
        "stream": False,
    }
    headers = {"Content-Type": "application/json; charset=utf-8"}
    last_error: BatchItemError | None = None

    for attempt_count in range(1, max_retries + 2):
        try:
            response = session.post(
                _endpoint(base_url),
                json=request_body,
                headers=headers,
                timeout=timeout,
            )
            return _parse_http_response(
                response,
                mentor_input,
                attempt_count,
                mode=mode,
                standard_tags_enabled=standard_tags_enabled,
                taxonomy_hash=taxonomy_hash,
            )
        except requests.Timeout:
            last_error = BatchItemError(
                "timeout",
                "local mentor service request timed out",
                attempt_count=attempt_count,
                retriable=True,
            )
        except requests.ConnectionError:
            last_error = BatchItemError(
                "connection_error",
                "could not connect to the local mentor service",
                attempt_count=attempt_count,
                retriable=True,
            )
        except requests.RequestException:
            last_error = BatchItemError(
                "http_request_failed",
                "local mentor service request failed",
                attempt_count=attempt_count,
                retriable=True,
            )
        except BatchItemError as exc:
            last_error = exc
        except Exception:
            last_error = BatchItemError(
                "unexpected_error",
                "unexpected error during local mentor service request",
                attempt_count=attempt_count,
                retriable=False,
            )

        if not last_error.retriable or attempt_count > max_retries:
            raise last_error
        if retry_sleep:
            sleep_fn(retry_sleep)

    raise RuntimeError("unreachable request retry state")


def request_simple_mentor_results_batch(
    mentor_inputs: list[MentorInput],
    *,
    base_url: str,
    timeout: float,
    max_retries: int,
    retry_sleep: float,
    session: Any,
    standard_tags_enabled: bool = False,
    taxonomy_hash: str | None = None,
    sleep_fn: Callable[[float], None] = time.sleep,
) -> list[SimpleMentorResult]:
    """Call the simple batch endpoint with bounded retries."""

    request_body = {
        "messages": [
            {
                "role": "user",
                "content": json.dumps(
                    {
                        "task": "extract_mentor_keywords_batch_simple",
                        "records": [
                            mentor_input.model_dump(by_alias=True, mode="json")
                            for mentor_input in mentor_inputs
                        ],
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                ),
            }
        ],
        "stream": False,
    }
    headers = {"Content-Type": "application/json; charset=utf-8"}
    last_error: BatchItemError | None = None

    for attempt_count in range(1, max_retries + 2):
        try:
            response = session.post(
                _endpoint(base_url),
                json=request_body,
                headers=headers,
                timeout=timeout,
            )
            return _parse_http_batch_response(
                response,
                mentor_inputs,
                attempt_count,
                standard_tags_enabled=standard_tags_enabled,
                taxonomy_hash=taxonomy_hash,
            )
        except requests.Timeout:
            last_error = BatchItemError(
                "timeout",
                "local mentor service batch request timed out",
                attempt_count=attempt_count,
                retriable=True,
                stage="http_extract_batch",
            )
        except requests.ConnectionError:
            last_error = BatchItemError(
                "connection_error",
                "could not connect to the local mentor service",
                attempt_count=attempt_count,
                retriable=True,
                stage="http_extract_batch",
            )
        except requests.RequestException:
            last_error = BatchItemError(
                "http_request_failed",
                "local mentor service batch request failed",
                attempt_count=attempt_count,
                retriable=True,
                stage="http_extract_batch",
            )
        except BatchItemError as exc:
            last_error = exc
            last_error.stage = "http_extract_batch"
        except Exception:
            last_error = BatchItemError(
                "unexpected_error",
                "unexpected error during local mentor service batch request",
                attempt_count=attempt_count,
                retriable=False,
                stage="http_extract_batch",
            )

        if not last_error.retriable or attempt_count > max_retries:
            raise last_error
        if retry_sleep:
            sleep_fn(retry_sleep)

    raise RuntimeError("unreachable batch request retry state")


def load_success_checkpoint(
    path: Path,
    *,
    mode: ExtractionMode,
    standard_tags_enabled: bool = False,
    taxonomy_hash: str | None = None,
    warn: Callable[[str], None] = print,
) -> dict[tuple[str, str], ExtractionResult]:
    successes: dict[tuple[str, str], ExtractionResult] = {}
    if not path.exists():
        return successes

    result_model = _result_model_for_mode(mode)
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                result = result_model.model_validate_json(line)
            except (ValidationError, ValueError):
                warn(
                    f"warning: skipped invalid or mismatched {mode} checkpoint line {line_number}"
                )
                continue
            if isinstance(result, SimpleMentorResult):
                if result.standard_tags_enabled != standard_tags_enabled:
                    warn(
                        "warning: skipped simple checkpoint line "
                        f"{line_number} due to standard-tags mode mismatch"
                    )
                    continue
                if (
                    standard_tags_enabled
                    and result.taxonomy_hash != taxonomy_hash
                ):
                    warn(
                        "warning: skipped simple checkpoint line "
                        f"{line_number} due to taxonomy hash mismatch"
                    )
                    continue
            successes[(result.mentor_id, result.record_hash)] = result
    return successes


def _append_json_line(handle: Any, payload: str | Mapping[str, Any]) -> None:
    if isinstance(payload, str):
        encoded = payload
    else:
        encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    handle.write(encoded + "\n")
    handle.flush()


def _failure_record(
    mentor_input: MentorInput,
    error: BatchItemError,
) -> dict[str, Any]:
    return {
        "mentor_id": mentor_input.mentor_id,
        "record_hash": mentor_input.record_hash,
        "source_file": mentor_input.source_file,
        "source_sheet": mentor_input.source_sheet,
        "source_row": mentor_input.source_row,
        "error_type": error.error_type,
        "sanitized_error_message": error.safe_message,
        "attempt_count": error.attempt_count,
        "failed_at": datetime.now(timezone.utc).isoformat(),
        "retriable": error.retriable,
        "stage": error.stage,
    }


def _excel_failure_records(report: ExcelLoadReport) -> list[dict[str, Any]]:
    records = []
    for row_issue in report.row_issues:
        if row_issue.issue.severity.value != "error":
            continue
        records.append(
            {
                "mentor_id": None,
                "record_hash": None,
                "source_file": report.summary.source_file,
                "source_sheet": report.summary.sheet,
                "source_row": row_issue.source_row,
                "error_type": "excel_invalid_row",
                "sanitized_error_message": sanitize_error_message(
                    row_issue.issue.message
                ),
                "attempt_count": 0,
                "failed_at": datetime.now(timezone.utc).isoformat(),
                "retriable": False,
                "stage": "excel_preprocessing",
            }
        )
    return records


def _write_final_results(path: Path, results: Iterable[ExtractionResult]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for result in results:
            handle.write(result.model_dump_json(by_alias=True) + "\n")


def _read_failure_history(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    if not path.exists():
        return records
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            try:
                payload = json.loads(line)
            except json.JSONDecodeError:
                continue
            if isinstance(payload, dict):
                payload["sanitized_error_message"] = sanitize_error_message(
                    str(payload.get("sanitized_error_message", ""))
                )
                records.append(payload)
    return records


def _evidence_rows(
    result: MentorResult,
    item: Any,
    raw_value: Any,
    normalized_value: Any,
    mapping_status: Any,
    extra: Mapping[str, Any] | None = None,
) -> list[list[Any]]:
    rows = []
    for evidence in item.evidence:
        row = [
            result.mentor_id,
            result.original_fields.mentor_name,
            raw_value,
            normalized_value,
            _enum_value(mapping_status),
            _enum_value(evidence.source_field),
            evidence.quote,
            _enum_value(evidence.match_type),
            _enum_value(item.confidence),
        ]
        if extra:
            row.extend(extra.values())
        rows.append(row)
    return rows


def _style_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    if sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_index in range(1, sheet.max_column + 1):
        values = [
            "" if sheet.cell(row, column_index).value is None else str(sheet.cell(row, column_index).value)
            for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        width = min(max((len(value) for value in values), default=10) + 2, 50)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)


def _add_sheet(
    workbook: Workbook,
    name: str,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    _style_sheet(sheet)


def write_review_excel(
    path: Path,
    results: list[MentorResult],
    failures: list[dict[str, Any]],
) -> None:
    """Create a new review workbook; never modify the source workbook."""

    workbook = Workbook()
    workbook.remove(workbook.active)

    overview_rows = []
    for result in results:
        original = result.original_fields
        overview_rows.append(
            [
                result.mentor_id,
                original.mentor_name,
                original.gender,
                original.city,
                original.career_years,
                original.coachable_levels,
                original.industry_tags,
                result.summary.value if result.summary else None,
                len(result.industry_tags),
                len(result.career_experiences),
                len(result.skills),
                len(result.credentials_and_awards),
                len(result.education),
                len(result.quality_issues),
                result.processing.status,
                result.processing.model_name,
                result.processing.latency_ms,
            ]
        )
    _add_sheet(
        workbook,
        "导师总览",
        [
            "mentor_id",
            "导师姓名",
            "性别",
            "城市",
            "职业年限",
            "可辅导学员职级",
            "原行业标签",
            "summary",
            "industry_count",
            "career_experience_count",
            "skill_count",
            "credential_count",
            "education_count",
            "quality_issue_count",
            "processing_status",
            "model_name",
            "latency_ms",
        ],
        overview_rows,
    )

    common_headers = [
        "mentor_id",
        "导师姓名",
        "raw",
        "normalized",
        "mapping_status",
        "evidence_source_field",
        "evidence_quote",
        "evidence_match_type",
        "confidence",
    ]

    industry_rows = []
    career_rows = []
    skill_rows = []
    credential_rows = []
    education_rows = []
    target_rows = []
    highlight_rows = []
    quality_rows = []

    for result in results:
        for item in result.industry_tags:
            industry_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.raw_industry,
                    item.normalized_industry,
                    item.mapping_status,
                    {"industry_category": item.industry_category},
                )
            )
        for item in result.career_experiences:
            career_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.raw_name,
                    item.normalized_name,
                    item.mapping_status,
                    {
                        "relationship": item.relationship.value,
                        "raw_title": item.raw_title,
                        "normalized_title": item.normalized_title,
                        "employment_status": item.employment_status.value,
                    },
                )
            )
        for item in result.skills:
            skill_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.raw_skill,
                    item.normalized_skill,
                    item.mapping_status,
                    {"skill_category": _enum_value(item.skill_category)},
                )
            )
        for item in result.credentials_and_awards:
            credential_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.raw_name,
                    item.normalized_name,
                    item.mapping_status,
                    {
                        "category": item.category.value,
                        "credential_status": item.credential_status.value,
                        "issuer_raw": item.issuer_raw,
                    },
                )
            )
        for item in result.education:
            education_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.institution_raw,
                    item.institution_normalized,
                    item.institution_mapping_status,
                    {
                        "degree_raw": item.degree_raw,
                        "degree_normalized": item.degree_normalized,
                        "major_raw": item.major_raw,
                        "major_normalized": item.major_normalized,
                    },
                )
            )
        for item in result.target_mentees:
            target_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.raw_audience,
                    item.normalized_audience,
                    item.mapping_status,
                    {"dimension": _enum_value(item.dimension)},
                )
            )
        for item in result.career_highlights:
            highlight_rows.extend(
                _evidence_rows(
                    result,
                    item,
                    item.statement,
                    item.statement,
                    None,
                    {"highlight_type": item.type},
                )
            )
        for issue in result.quality_issues:
            quality_rows.append(
                [
                    result.mentor_id,
                    result.original_fields.mentor_name,
                    issue.code,
                    issue.severity.value,
                    issue.path,
                    _enum_value(issue.source_field),
                    issue.message,
                ]
            )

    _add_sheet(
        workbook,
        "行业标签",
        common_headers + ["industry_category"],
        industry_rows,
    )
    _add_sheet(
        workbook,
        "任职经历",
        common_headers
        + ["relationship", "raw_title", "normalized_title", "employment_status"],
        career_rows,
    )
    _add_sheet(
        workbook,
        "技能",
        common_headers + ["skill_category"],
        skill_rows,
    )
    _add_sheet(
        workbook,
        "证书资质奖项",
        common_headers + ["category", "credential_status", "issuer_raw"],
        credential_rows,
    )
    _add_sheet(
        workbook,
        "教育背景",
        common_headers
        + ["degree_raw", "degree_normalized", "major_raw", "major_normalized"],
        education_rows,
    )
    _add_sheet(
        workbook,
        "辅导人群",
        common_headers + ["dimension"],
        target_rows,
    )
    _add_sheet(
        workbook,
        "职业亮点",
        common_headers + ["highlight_type"],
        highlight_rows,
    )
    _add_sheet(
        workbook,
        "质量问题",
        [
            "mentor_id",
            "导师姓名",
            "code",
            "severity",
            "path",
            "source_field",
            "message",
        ],
        quality_rows,
    )
    _add_sheet(
        workbook,
        "处理失败",
        [
            "mentor_id",
            "record_hash",
            "source_file",
            "source_sheet",
            "source_row",
            "error_type",
            "sanitized_error_message",
            "attempt_count",
            "failed_at",
            "retriable",
            "stage",
        ],
        (
            [
                failure.get("mentor_id"),
                failure.get("record_hash"),
                failure.get("source_file"),
                failure.get("source_sheet"),
                failure.get("source_row"),
                failure.get("error_type"),
                sanitize_error_message(
                    str(failure.get("sanitized_error_message", ""))
                ),
                failure.get("attempt_count"),
                failure.get("failed_at"),
                failure.get("retriable"),
                failure.get("stage"),
            ]
            for failure in failures
        ),
    )

    workbook.save(path)


def _join_keywords(values: list[str]) -> str:
    return "；".join(values)


def write_simple_review_excel(
    path: Path,
    results: list[SimpleMentorResult],
    failures: list[dict[str, Any]],
    *,
    standard_tags_enabled: bool | None = None,
) -> None:
    """Create a compact simple-mode review workbook."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    tagged = (
        any(result.standard_tags_enabled for result in results)
        if standard_tags_enabled is None
        else standard_tags_enabled
    )

    overview_rows = []
    detail_rows = []
    categories = [
        "industries",
        "companies",
        "roles",
        "skills",
        "credentials",
        "education",
        "target_mentees",
        "highlights",
        "keywords",
    ]
    if tagged:
        categories.append("raw_keywords")

    industry_tag_rows = []
    position_tag_rows = []
    company_tag_rows = []
    review_rows = []

    for result in results:
        original = result.original_fields
        extraction = result.extraction
        overview_row = [
            result.mentor_id,
            original.mentor_name,
            original.gender,
            original.city,
            original.career_years,
            original.industry_tags,
            _join_keywords(extraction.industries),
            _join_keywords(extraction.companies),
            _join_keywords(extraction.roles),
            _join_keywords(extraction.skills),
            _join_keywords(extraction.credentials),
            _join_keywords(extraction.education),
            _join_keywords(extraction.target_mentees),
            _join_keywords(extraction.highlights),
            _join_keywords(extraction.keywords),
            extraction.summary,
            result.processing.latency_ms,
            result.processing.attempt_count,
        ]
        if tagged:
            overview_row.extend(
                [
                    _join_keywords(
                        [item.tag for item in extraction.industry_tags]
                    ),
                    _join_keywords(
                        [item.tag for item in extraction.position_tags]
                    ),
                    extraction.review_required,
                    _join_keywords(extraction.review_reasons),
                    result.taxonomy_hash,
                ]
            )
        overview_rows.append(overview_row)
        for category in categories:
            for keyword in getattr(extraction, category):
                detail_rows.append(
                    [
                        result.mentor_id,
                        original.mentor_name,
                        category,
                        keyword,
                    ]
                )
        if tagged:
            for item in extraction.industry_tags:
                industry_tag_rows.append(
                    [
                        result.mentor_id,
                        original.mentor_name,
                        item.tag,
                        item.confidence,
                        item.evidence,
                    ]
                )
            for item in extraction.position_tags:
                position_tag_rows.append(
                    [
                        result.mentor_id,
                        original.mentor_name,
                        item.tag,
                        _enum_value(item.relation_type),
                        item.confidence,
                        _join_keywords(item.raw_keywords),
                        item.evidence,
                    ]
                )
            for item in extraction.company_tags:
                company_tag_rows.append(
                    [
                        result.mentor_id,
                        original.mentor_name,
                        item.company_name,
                        item.company_type,
                        item.industry_tag,
                        item.confidence,
                        item.evidence,
                    ]
                )
            for reason in extraction.review_reasons:
                review_rows.append(
                    [
                        result.mentor_id,
                        original.mentor_name,
                        extraction.review_required,
                        reason,
                    ]
                )

    overview_headers = [
        "mentor_id",
        "导师姓名",
        "性别",
        "城市",
        "职业年限",
        "行业标签原文",
        "industries",
        "companies",
        "roles",
        "skills",
        "credentials",
        "education",
        "target_mentees",
        "highlights",
        "keywords",
        "summary",
        "latency_ms",
        "attempt_count",
    ]
    if tagged:
        overview_headers.extend(
            [
                "standard_industry_tags",
                "standard_position_tags",
                "review_required",
                "review_reasons",
                "taxonomy_hash",
            ]
        )
    _add_sheet(
        workbook,
        "导师总览",
        overview_headers,
        overview_rows,
    )
    _add_sheet(
        workbook,
        "关键词明细",
        ["mentor_id", "导师姓名", "category", "keyword"],
        detail_rows,
    )
    if tagged:
        _add_sheet(
            workbook,
            "标准行业标签",
            ["mentor_id", "导师姓名", "tag", "confidence", "evidence"],
            industry_tag_rows,
        )
        _add_sheet(
            workbook,
            "标准职位标签",
            [
                "mentor_id",
                "导师姓名",
                "tag",
                "relation_type",
                "confidence",
                "raw_keywords",
                "evidence",
            ],
            position_tag_rows,
        )
        _add_sheet(
            workbook,
            "公司标签",
            [
                "mentor_id",
                "导师姓名",
                "company_name",
                "company_type",
                "industry_tag",
                "confidence",
                "evidence",
            ],
            company_tag_rows,
        )
        _add_sheet(
            workbook,
            "人工审核项",
            ["mentor_id", "导师姓名", "review_required", "review_reason"],
            review_rows,
        )
    _add_sheet(
        workbook,
        "处理失败",
        [
            "mentor_id",
            "record_hash",
            "source_file",
            "source_sheet",
            "source_row",
            "error_type",
            "sanitized_error_message",
            "attempt_count",
            "failed_at",
            "retriable",
            "stage",
        ],
        (
            [
                failure.get("mentor_id"),
                failure.get("record_hash"),
                failure.get("source_file"),
                failure.get("source_sheet"),
                failure.get("source_row"),
                failure.get("error_type"),
                sanitize_error_message(
                    str(failure.get("sanitized_error_message", ""))
                ),
                failure.get("attempt_count"),
                failure.get("failed_at"),
                failure.get("retriable"),
                failure.get("stage"),
            ]
            for failure in failures
        ),
    )

    workbook.save(path)


def _validate_config(config: BatchConfig) -> None:
    if config.resume and config.force:
        raise ValueError("resume and force are mutually exclusive")
    if config.limit is not None and config.limit < 0:
        raise ValueError("limit must be >= 0")
    if config.offset < 0:
        raise ValueError("offset must be >= 0")
    if config.timeout <= 0:
        raise ValueError("timeout must be > 0")
    if config.max_retries < 0:
        raise ValueError("max_retries must be >= 0")
    if config.retry_sleep < 0:
        raise ValueError("retry_sleep must be >= 0")
    if config.mode not in EXTRACTION_MODES:
        raise ValueError("mode must be 'simple' or 'full'")
    if config.batch_size <= 0:
        raise ValueError("batch_size must be > 0")


def _chunks(
    items: list[tuple[int, MentorInput]],
    size: int,
) -> Iterable[list[tuple[int, MentorInput]]]:
    for start in range(0, len(items), size):
        yield items[start : start + size]


def run_batch(
    config: BatchConfig,
    *,
    session: Any | None = None,
    sleep_fn: Callable[[float], None] = time.sleep,
    printer: Callable[[str], None] = print,
) -> BatchSummary:
    """Run one batch without allowing a single row failure to abort the batch."""

    _validate_config(config)
    standard_tags_enabled = (
        config.mode == "simple" and config.standard_tags_enabled
    )
    taxonomy_hash = (
        load_tag_taxonomy(config.taxonomy_path).taxonomy_hash
        if standard_tags_enabled
        else None
    )
    report = load_mentor_inputs_with_report(config.input_path)
    all_inputs = report.inputs
    selected = all_inputs[config.offset :]
    if config.limit is not None:
        selected = selected[: config.limit]

    if config.dry_run:
        summary = BatchSummary(
            source_file=report.summary.source_file,
            valid_mentor_count=len(all_inputs),
            selected_count=len(selected),
            success_count=0,
            failed_count=0,
            skipped_resume_count=0,
            final_result_count=0,
            dry_run=True,
        )
        printer(json.dumps(summary.to_dict(), ensure_ascii=False, sort_keys=True))
        return summary

    config.output_dir.mkdir(parents=True, exist_ok=True)
    checkpoint_path = config.output_dir / SUCCESS_CHECKPOINT_RELATIVE
    checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
    failed_path = config.output_dir / FAILED_ROWS_NAME
    final_path = config.output_dir / FINAL_RESULTS_NAME
    review_path = config.output_dir / REVIEW_EXCEL_NAME

    checkpoint_results = load_success_checkpoint(
        checkpoint_path,
        mode=config.mode,
        standard_tags_enabled=standard_tags_enabled,
        taxonomy_hash=taxonomy_hash,
        warn=printer,
    )
    failed_current_keys: set[tuple[str, str]] = set()
    excel_failure_records = _excel_failure_records(report)
    success_count = 0
    failed_count = len(excel_failure_records)
    skipped_resume_count = 0
    http_session = session or requests.Session()

    with checkpoint_path.open("a", encoding="utf-8", newline="\n") as success_handle, failed_path.open(
        "a", encoding="utf-8", newline="\n"
    ) as failed_handle:
        for record in excel_failure_records:
            _append_json_line(failed_handle, record)

        total_selected = len(selected)
        pending: list[tuple[int, MentorInput]] = []
        for index, mentor_input in enumerate(selected, start=1):
            key = (mentor_input.mentor_id, mentor_input.record_hash)
            if config.resume and key in checkpoint_results:
                skipped_resume_count += 1
                printer(
                    f"[{index}/{total_selected}] {mentor_input.mentor_id} skipped resume"
                )
                continue
            pending.append((index, mentor_input))

        def process_single(index: int, mentor_input: MentorInput) -> None:
            nonlocal success_count, failed_count
            key = (mentor_input.mentor_id, mentor_input.record_hash)
            try:
                result = request_mentor_result(
                    mentor_input,
                    base_url=config.base_url,
                    timeout=config.timeout,
                    max_retries=config.max_retries,
                    retry_sleep=config.retry_sleep,
                    session=http_session,
                    mode=config.mode,
                    standard_tags_enabled=standard_tags_enabled,
                    taxonomy_hash=taxonomy_hash,
                    sleep_fn=sleep_fn,
                )
            except BatchItemError as exc:
                failed_count += 1
                failed_current_keys.add(key)
                _append_json_line(failed_handle, _failure_record(mentor_input, exc))
                printer(
                    f"[{index}/{total_selected}] {mentor_input.mentor_id} "
                    f"failed error_type={exc.error_type}"
                )
                return
            except Exception:
                failed_count += 1
                failed_current_keys.add(key)
                error = BatchItemError(
                    "unexpected_error",
                    "unexpected error while processing mentor row",
                    attempt_count=0,
                    retriable=False,
                    stage="batch_processing",
                )
                _append_json_line(failed_handle, _failure_record(mentor_input, error))
                printer(
                    f"[{index}/{total_selected}] {mentor_input.mentor_id} "
                    "failed error_type=unexpected_error"
                )
                return

            success_count += 1
            checkpoint_results[key] = result
            _append_json_line(
                success_handle,
                result.model_dump_json(by_alias=True),
            )
            printer(
                f"[{index}/{total_selected}] {mentor_input.mentor_id} "
                f"success latency={result.processing.latency_ms}"
            )

        def record_batch_success(
            index: int,
            mentor_input: MentorInput,
            result: SimpleMentorResult,
        ) -> None:
            nonlocal success_count
            key = (mentor_input.mentor_id, mentor_input.record_hash)
            success_count += 1
            checkpoint_results[key] = result
            _append_json_line(
                success_handle,
                result.model_dump_json(by_alias=True),
            )
            printer(
                f"[{index}/{total_selected}] {mentor_input.mentor_id} "
                f"success latency={result.processing.latency_ms}"
            )

        if config.mode == "simple" and config.batch_size > 1:
            for group in _chunks(pending, config.batch_size):
                group_inputs = [mentor_input for _, mentor_input in group]
                try:
                    batch_results = request_simple_mentor_results_batch(
                        group_inputs,
                        base_url=config.base_url,
                        timeout=config.timeout,
                        max_retries=config.max_retries,
                        retry_sleep=config.retry_sleep,
                        session=http_session,
                        standard_tags_enabled=standard_tags_enabled,
                        taxonomy_hash=taxonomy_hash,
                        sleep_fn=sleep_fn,
                    )
                    results_by_id = {
                        result.mentor_id: result for result in batch_results
                    }
                    for index, mentor_input in group:
                        record_batch_success(
                            index,
                            mentor_input,
                            results_by_id[mentor_input.mentor_id],
                        )
                except BatchItemError as exc:
                    printer(
                        f"batch fallback size={len(group)} "
                        f"error_type={exc.error_type}"
                    )
                    for index, mentor_input in group:
                        process_single(index, mentor_input)
                except Exception:
                    printer(
                        f"batch fallback size={len(group)} "
                        "error_type=unexpected_error"
                    )
                    for index, mentor_input in group:
                        process_single(index, mentor_input)
        else:
            for index, mentor_input in pending:
                process_single(index, mentor_input)

    final_results: list[ExtractionResult] = []
    for mentor_input in all_inputs:
        key = (mentor_input.mentor_id, mentor_input.record_hash)
        if key in failed_current_keys:
            continue
        result = checkpoint_results.get(key)
        if result is not None:
            final_results.append(result)

    _write_final_results(final_path, final_results)
    if not config.no_excel:
        if config.mode == "simple":
            write_simple_review_excel(
                review_path,
                [
                    result
                    for result in final_results
                    if isinstance(result, SimpleMentorResult)
                ],
                _read_failure_history(failed_path),
                standard_tags_enabled=standard_tags_enabled,
            )
        else:
            write_review_excel(
                review_path,
                [
                    result
                    for result in final_results
                    if isinstance(result, MentorResult)
                ],
                _read_failure_history(failed_path),
            )

    summary = BatchSummary(
        source_file=report.summary.source_file,
        valid_mentor_count=len(all_inputs),
        selected_count=len(selected),
        success_count=success_count,
        failed_count=failed_count,
        skipped_resume_count=skipped_resume_count,
        final_result_count=len(final_results),
        dry_run=False,
    )
    if config.verbose:
        printer(json.dumps(summary.to_dict(), ensure_ascii=False, sort_keys=True))
    return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Batch mentor extraction through the local AgentRun endpoint."
    )
    parser.add_argument("--input", "-i", required=True, type=Path)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--limit", type=_non_negative_int)
    parser.add_argument("--offset", type=_non_negative_int, default=0)
    parser.add_argument(
        "--batch-size",
        type=_positive_int,
        default=10,
        help="Simple-mode HTTP batch size; use 1 for single-record requests.",
    )
    parser.add_argument(
        "--mode",
        choices=EXTRACTION_MODES,
        default="simple",
        help="Result schema to expect from the local service.",
    )
    mode_group = parser.add_mutually_exclusive_group()
    mode_group.add_argument("--resume", action="store_true")
    mode_group.add_argument("--force", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--timeout", type=_positive_float, default=120.0)
    parser.add_argument("--max-retries", type=_non_negative_int, default=2)
    parser.add_argument("--retry-sleep", type=_non_negative_float, default=1.0)
    parser.add_argument("--no-excel", action="store_true")
    parser.add_argument("--verbose", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    config = BatchConfig(
        input_path=args.input,
        output_dir=args.output_dir,
        base_url=args.base_url,
        limit=args.limit,
        offset=args.offset,
        resume=args.resume,
        force=args.force,
        dry_run=args.dry_run,
        timeout=args.timeout,
        max_retries=args.max_retries,
        retry_sleep=args.retry_sleep,
        no_excel=args.no_excel,
        verbose=args.verbose,
        mode=args.mode,
        batch_size=args.batch_size,
        standard_tags_enabled=_boolean_env("ENABLE_STANDARD_TAGS", False),
        taxonomy_path=Path(
            os.getenv("TAG_TAXONOMY_PATH", "configs/职位类型_2.txt")
        ),
    )
    summary = run_batch(config)
    if not config.verbose and not config.dry_run:
        print(json.dumps(summary.to_dict(), ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
