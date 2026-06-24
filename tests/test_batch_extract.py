"""Offline Stage 5 tests for batch processing, resume, and exports."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import pandas as pd
import pytest
import requests
from openpyxl import load_workbook

from batch_extract import (
    FAILED_ROWS_NAME,
    FINAL_RESULTS_NAME,
    REVIEW_EXCEL_NAME,
    REVIEW_SHEETS,
    SUCCESS_CHECKPOINT_RELATIVE,
    BatchConfig,
    build_parser,
    run_batch,
)
from mentor_agent.excel_io import EXPECTED_SOURCE_FIELDS
from mentor_agent.schemas import MentorInput, MentorResult


def _mentor_row(sequence_no: int, city: str = "杭州") -> dict[str, Any]:
    return {
        "序号": sequence_no,
        "导师姓名": f"匿名导师{sequence_no}",
        "性别": "女",
        "城市": city,
        "职业年限": 8,
        "可辅导学员职级": "P5-P7",
        "行业标签": "软件与信息技术",
        "从业经历": f"曾任匿名公司{sequence_no}产品负责人。",
        "背景经验": "擅长产品规划与职业辅导。",
    }


def _write_source(path: Path, rows: list[dict[str, Any]]) -> None:
    pd.DataFrame(rows, columns=EXPECTED_SOURCE_FIELDS).to_excel(
        path,
        sheet_name="服务导师",
        index=False,
        engine="openpyxl",
    )


def _result_for_input(mentor_input: MentorInput) -> MentorResult:
    original = mentor_input.original_fields
    return MentorResult.model_validate(
        {
            "schema_version": "1.0",
            "prompt_version": "mentor-extraction-v1.1",
            "industry_taxonomy_version": "industry-soft-v1",
            "mentor_id": mentor_input.mentor_id,
            "record_hash": mentor_input.record_hash,
            "source": {
                "file": mentor_input.source_file,
                "sheet": mentor_input.source_sheet,
                "row": mentor_input.source_row,
            },
            "original_fields": original.model_dump(by_alias=True),
            "normalized_profile": {
                "name": {
                    "raw_value": original.mentor_name,
                    "normalized_value": original.mentor_name,
                    "mapping_status": "mapped",
                }
            },
            "skills": [
                {
                    "raw_skill": "产品规划",
                    "normalized_skill": "产品规划",
                    "mapping_status": "unmapped",
                    "evidence": [
                        {
                            "source_field": "背景经验",
                            "quote": "产品规划",
                            "match_type": "strict",
                        }
                    ],
                    "confidence": "high",
                }
            ],
            "career_experiences": [
                {
                    "raw_name": f"匿名公司{original.sequence_no}",
                    "normalized_name": f"匿名公司{original.sequence_no}",
                    "mapping_status": "unmapped",
                    "relationship": "employer",
                    "employment_status": "former",
                    "evidence": [
                        {
                            "source_field": "从业经历",
                            "quote": f"曾任匿名公司{original.sequence_no}产品负责人",
                            "match_type": "strict",
                        }
                    ],
                    "confidence": "high",
                }
            ],
            "quality_issues": [],
            "processing": {
                "status": "success",
                "attempt_count": 1,
                "processed_at": datetime.now(timezone.utc).isoformat(),
                "model_service_name": "fake-http",
                "model_name": "fake-model",
                "latency_ms": 7,
            },
        }
    )


class FakeResponse:
    def __init__(
        self,
        status_code: int = 200,
        payload: Any = None,
        *,
        json_error: bool = False,
    ) -> None:
        self.status_code = status_code
        self.payload = payload
        self.json_error = json_error

    def json(self) -> Any:
        if self.json_error:
            raise ValueError("not json")
        return self.payload


def _success_response(request_kwargs: dict[str, Any]) -> FakeResponse:
    content = request_kwargs["json"]["messages"][0]["content"]
    mentor_input = MentorInput.model_validate_json(content)
    result = _result_for_input(mentor_input)
    return FakeResponse(
        payload={
            "choices": [
                {"message": {"content": result.model_dump_json(by_alias=True)}}
            ]
        }
    )


class FakeSession:
    def __init__(self, actions: list[Any] | None = None) -> None:
        self.actions = list(actions or [])
        self.calls: list[dict[str, Any]] = []

    def post(self, url: str, **kwargs: Any) -> FakeResponse:
        call = {"url": url, **kwargs}
        self.calls.append(call)
        action = self.actions.pop(0) if self.actions else _success_response
        if isinstance(action, Exception):
            raise action
        if callable(action):
            return action(call)
        return action


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


def test_dry_run_does_not_call_http_or_write_outputs(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1), _mentor_row(2)])
    session = FakeSession()

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            dry_run=True,
            limit=1,
        ),
        session=session,
        printer=lambda message: None,
    )

    assert summary.dry_run is True
    assert summary.selected_count == 1
    assert session.calls == []
    assert not output.exists()


def test_limit_offset_success_checkpoint_final_jsonl_and_review_excel(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    rows = [
        _mentor_row(1),
        {"序号": 99},
        _mentor_row(2),
        _mentor_row(3),
        _mentor_row(4),
    ]
    _write_source(source, rows)
    source_digest = hashlib.sha256(source.read_bytes()).hexdigest()
    session = FakeSession()

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            offset=1,
            limit=2,
            max_retries=0,
        ),
        session=session,
        printer=lambda message: None,
    )

    assert summary.success_count == 2
    assert [
        MentorInput.model_validate_json(call["json"]["messages"][0]["content"]).mentor_id
        for call in session.calls
    ] == ["service_mentor:2", "service_mentor:3"]
    assert all(call["json"]["stream"] is False for call in session.calls)
    assert len(_read_jsonl(output / SUCCESS_CHECKPOINT_RELATIVE)) == 2
    assert len(_read_jsonl(output / FINAL_RESULTS_NAME)) == 2
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_digest

    workbook = load_workbook(output / REVIEW_EXCEL_NAME, read_only=True)
    assert tuple(workbook.sheetnames) == REVIEW_SHEETS
    assert workbook["导师总览"].max_row == 3
    assert workbook["技能"].max_row == 3
    assert workbook["任职经历"].max_row == 3


def test_failure_is_written_and_does_not_stop_later_records(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1), _mentor_row(2)])
    session = FakeSession(
        [
            FakeResponse(status_code=500, payload={"error": "secret body"}),
            _success_response,
        ]
    )

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=lambda message: None,
    )

    failures = _read_jsonl(output / FAILED_ROWS_NAME)
    results = _read_jsonl(output / FINAL_RESULTS_NAME)
    assert summary.failed_count == 1
    assert summary.success_count == 1
    assert failures[0]["mentor_id"] == "service_mentor:1"
    assert failures[0]["error_type"] == "http_status_error"
    assert results[0]["mentor_id"] == "service_mentor:2"


def test_retriable_failure_can_succeed_on_next_attempt(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])
    session = FakeSession(
        [FakeResponse(status_code=500, payload={}), _success_response]
    )
    sleeps: list[float] = []

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            max_retries=1,
            retry_sleep=0.25,
            no_excel=True,
        ),
        session=session,
        sleep_fn=sleeps.append,
        printer=lambda message: None,
    )

    assert summary.success_count == 1
    assert summary.failed_count == 0
    assert len(session.calls) == 2
    assert sleeps == [0.25]


def test_resume_skips_matching_success(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])

    run_batch(
        BatchConfig(input_path=source, output_dir=output, max_retries=0),
        session=FakeSession(),
        printer=lambda message: None,
    )
    resume_session = FakeSession()

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            resume=True,
            max_retries=0,
        ),
        session=resume_session,
        printer=lambda message: None,
    )

    assert summary.skipped_resume_count == 1
    assert summary.final_result_count == 1
    assert resume_session.calls == []


def test_changed_record_hash_is_not_skipped_by_resume(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1, city="杭州")])
    run_batch(
        BatchConfig(input_path=source, output_dir=output, max_retries=0),
        session=FakeSession(),
        printer=lambda message: None,
    )
    first_result = _read_jsonl(output / FINAL_RESULTS_NAME)[0]

    _write_source(source, [_mentor_row(1, city="上海")])
    changed_session = FakeSession()
    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            resume=True,
            max_retries=0,
        ),
        session=changed_session,
        printer=lambda message: None,
    )
    current_result = _read_jsonl(output / FINAL_RESULTS_NAME)[0]

    assert summary.skipped_resume_count == 0
    assert len(changed_session.calls) == 1
    assert current_result["record_hash"] != first_result["record_hash"]
    assert len(_read_jsonl(output / SUCCESS_CHECKPOINT_RELATIVE)) == 2


def test_force_failure_does_not_reuse_old_success_for_current_hash(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])
    run_batch(
        BatchConfig(input_path=source, output_dir=output, max_retries=0),
        session=FakeSession(),
        printer=lambda message: None,
    )

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            force=True,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([FakeResponse(status_code=500, payload={})]),
        printer=lambda message: None,
    )

    assert summary.failed_count == 1
    assert summary.final_result_count == 0
    assert _read_jsonl(output / FINAL_RESULTS_NAME) == []


def test_corrupt_checkpoint_line_warns_and_continues(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    checkpoint = output / SUCCESS_CHECKPOINT_RELATIVE
    checkpoint.parent.mkdir(parents=True)
    checkpoint.write_text("{broken checkpoint\n", encoding="utf-8")
    _write_source(source, [_mentor_row(1)])
    messages: list[str] = []
    session = FakeSession()

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            resume=True,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=messages.append,
    )

    assert summary.success_count == 1
    assert len(session.calls) == 1
    assert any("invalid checkpoint line 1" in message for message in messages)
    assert all("broken checkpoint" not in message for message in messages)


@pytest.mark.parametrize(
    ("response", "expected_error_type"),
    [
        (FakeResponse(status_code=500, payload={}), "http_status_error"),
        (FakeResponse(payload={"choices": []}), "missing_message_content"),
        (
            FakeResponse(payload={"choices": [{"message": {"content": ""}}]}),
            "missing_message_content",
        ),
        (
            FakeResponse(
                payload={"choices": [{"message": {"content": "not-json"}}]}
            ),
            "invalid_mentor_result_json",
        ),
        (
            FakeResponse(payload={"choices": [{"message": {"content": "{}"}}]}),
            "mentor_result_validation_failed",
        ),
        (FakeResponse(payload={"error": {"message": "server error"}}), "invalid_response_shape"),
        (FakeResponse(json_error=True), "invalid_response_shape"),
    ],
)
def test_http_and_response_failures_are_classified(
    tmp_path: Path,
    response: FakeResponse,
    expected_error_type: str,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([response]),
        printer=lambda message: None,
    )

    failure = _read_jsonl(output / FAILED_ROWS_NAME)[0]
    assert summary.failed_count == 1
    assert failure["error_type"] == expected_error_type
    assert _read_jsonl(output / FINAL_RESULTS_NAME) == []


def test_sensitive_request_error_is_not_written_or_printed(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])
    printed: list[str] = []
    session = FakeSession(
        [requests.ConnectionError("Authorization Bearer secret-token .env API Key=x")]
    )

    run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=printed.append,
    )

    combined = (output / FAILED_ROWS_NAME).read_text(encoding="utf-8") + repr(printed)
    for secret in ("secret-token", "Authorization", "API Key", ".env"):
        assert secret not in combined


def test_resume_and_force_are_mutually_exclusive() -> None:
    parser = build_parser()

    with pytest.raises(SystemExit):
        parser.parse_args(
            ["--input", "source.xlsx", "--resume", "--force"]
        )
