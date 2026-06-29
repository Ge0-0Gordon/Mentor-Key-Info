"""Batch tests for simple mode and mode mismatch handling."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from batch_extract import (
    FAILED_ROWS_NAME,
    FINAL_RESULTS_NAME,
    REVIEW_EXCEL_NAME,
    SIMPLE_REVIEW_SHEETS,
    SUCCESS_CHECKPOINT_RELATIVE,
    TAGGED_SIMPLE_REVIEW_SHEETS,
    BatchConfig,
    load_success_checkpoint,
    run_batch,
)
from mentor_agent.schemas import MentorInput
from mentor_agent.simple_schemas import SimpleMentorBatchResult, SimpleMentorResult
from mentor_agent.tag_taxonomy import load_tag_taxonomy
from tests.test_batch_extract import (
    FakeResponse,
    FakeSession,
    _mentor_row,
    _read_jsonl,
    _success_response,
    _write_source,
)


def _simple_result_for_input(mentor_input: MentorInput) -> SimpleMentorResult:
    original = mentor_input.original_fields
    return SimpleMentorResult.model_validate(
        {
            "schema_version": "simple-v1",
            "prompt_version": "mentor-simple-v1",
            "mentor_id": mentor_input.mentor_id,
            "record_hash": mentor_input.record_hash,
            "source": {
                "file": mentor_input.source_file,
                "sheet": mentor_input.source_sheet,
                "row": mentor_input.source_row,
            },
            "original_fields": original.model_dump(by_alias=True),
            "extraction": {
                "industries": ["software"],
                "companies": [f"anonymous-company-{original.sequence_no}"],
                "roles": ["product owner"],
                "skills": ["product planning", "career coaching"],
                "credentials": [],
                "education": [],
                "target_mentees": ["P5-P7"],
                "highlights": ["anonymous project experience"],
                "keywords": ["product", "coaching"],
                "summary": "Synthetic simple result for offline tests.",
            },
            "processing": {
                "status": "success",
                "attempt_count": 1,
                "processed_at": datetime.now(timezone.utc).isoformat(),
                "model_service_name": "fake-http",
                "model_name": "fake-simple-model",
                "latency_ms": 5,
            },
        }
    )


def _simple_success_response(request_kwargs: dict[str, Any]) -> FakeResponse:
    content = request_kwargs["json"]["messages"][0]["content"]
    mentor_input = MentorInput.model_validate_json(content)
    result = _simple_result_for_input(mentor_input)
    return FakeResponse(
        payload={
            "choices": [
                {"message": {"content": result.model_dump_json(by_alias=True)}}
            ]
        }
    )


def _simple_batch_success_response(request_kwargs: dict[str, Any]) -> FakeResponse:
    content = request_kwargs["json"]["messages"][0]["content"]
    payload = json.loads(content)
    results = [
        _simple_result_for_input(MentorInput.model_validate(record))
        for record in payload["records"]
    ]
    batch_result = SimpleMentorBatchResult(
        schema_version="simple-batch-v1",
        results=results,
    )
    return FakeResponse(
        payload={
            "choices": [
                {"message": {"content": batch_result.model_dump_json(by_alias=True)}}
            ]
        }
    )


def _tagged_result_for_input(
    mentor_input: MentorInput,
    taxonomy_hash: str,
) -> SimpleMentorResult:
    payload = _simple_result_for_input(mentor_input).model_dump(
        by_alias=True,
        mode="json",
    )
    payload["prompt_version"] = "mentor-simple-tagged-v1"
    payload["standard_tags_enabled"] = True
    payload["taxonomy_hash"] = taxonomy_hash
    payload["extraction"].update(
        {
            "industry_tags": [
                {
                    "tag": "AI/互联网/IT",
                    "confidence": 0.91,
                    "evidence": "软件与信息技术",
                }
            ],
            "position_tags": [
                {
                    "tag": "产品经理",
                    "relation_type": "firsthand_role",
                    "confidence": 0.93,
                    "raw_keywords": ["产品负责人"],
                    "evidence": "产品负责人",
                }
            ],
            "company_tags": [
                {
                    "company_name": f"匿名公司{mentor_input.original_fields.sequence_no}",
                    "company_type": "互联网公司",
                    "industry_tag": "AI/互联网/IT",
                    "confidence": 0.9,
                    "evidence": f"匿名公司{mentor_input.original_fields.sequence_no}",
                }
            ],
            "raw_keywords": ["产品负责人"],
            "review_required": False,
            "review_reasons": [],
        }
    )
    return SimpleMentorResult.model_validate(payload)


def _tagged_batch_success_response(
    taxonomy_hash: str,
):
    def response(request_kwargs: dict[str, Any]) -> FakeResponse:
        content = request_kwargs["json"]["messages"][0]["content"]
        payload = json.loads(content)
        results = [
            _tagged_result_for_input(
                MentorInput.model_validate(record),
                taxonomy_hash,
            )
            for record in payload["records"]
        ]
        batch_result = SimpleMentorBatchResult(
            schema_version="simple-batch-v1",
            results=results,
        )
        return FakeResponse(
            payload={
                "choices": [
                    {
                        "message": {
                            "content": batch_result.model_dump_json(
                                by_alias=True
                            )
                        }
                    }
                ]
            }
        )

    return response


def test_simple_mode_writes_jsonl_and_three_sheet_review_excel(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1), _mentor_row(2), _mentor_row(3)])

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=3,
            max_retries=0,
        ),
        session=FakeSession([_simple_batch_success_response]),
        printer=lambda message: None,
    )

    assert summary.success_count == 3
    checkpoint_rows = _read_jsonl(output / SUCCESS_CHECKPOINT_RELATIVE)
    final_rows = _read_jsonl(output / FINAL_RESULTS_NAME)
    assert len(checkpoint_rows) == 3
    assert len(final_rows) == 3
    SimpleMentorResult.model_validate(final_rows[0])

    workbook = load_workbook(output / REVIEW_EXCEL_NAME, read_only=True)
    assert tuple(workbook.sheetnames) == SIMPLE_REVIEW_SHEETS
    assert workbook[workbook.sheetnames[0]].max_row == 4
    assert workbook[workbook.sheetnames[1]].max_row > 3


def test_tagged_mode_writes_seven_sheet_review_excel(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    taxonomy_path = Path("configs/职位类型_2.txt")
    taxonomy_hash = load_tag_taxonomy(taxonomy_path).taxonomy_hash
    _write_source(source, [_mentor_row(1)])

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=2,
            max_retries=0,
            standard_tags_enabled=True,
            taxonomy_path=taxonomy_path,
        ),
        session=FakeSession([_tagged_batch_success_response(taxonomy_hash)]),
        printer=lambda message: None,
    )

    assert summary.success_count == 1
    result = SimpleMentorResult.model_validate(
        _read_jsonl(output / FINAL_RESULTS_NAME)[0]
    )
    assert result.standard_tags_enabled is True
    assert result.taxonomy_hash == taxonomy_hash
    workbook = load_workbook(output / REVIEW_EXCEL_NAME, read_only=True)
    assert tuple(workbook.sheetnames) == TAGGED_SIMPLE_REVIEW_SHEETS
    assert workbook["标准行业标签"].max_row == 2
    assert workbook["标准职位标签"].max_row == 2
    assert workbook["公司标签"].max_row == 2
    assert workbook["人工审核项"].max_row == 1


def test_checkpoint_requires_matching_taxonomy_hash(tmp_path: Path) -> None:
    checkpoint = tmp_path / "checkpoint.jsonl"
    mentor_input = MentorInput.model_validate(
        {
            "task": "extract_mentor_key_info",
            "input_schema_version": "1.0",
            "mentor_id": "service_mentor:1",
            "record_hash": "1" * 64,
            "source_file": "source.xlsx",
            "source_sheet": "服务导师",
            "source_row": 2,
            "original_fields": _mentor_row(1),
        }
    )
    result = _tagged_result_for_input(mentor_input, "a" * 64)
    checkpoint.write_text(result.model_dump_json() + "\n", encoding="utf-8")

    matching = load_success_checkpoint(
        checkpoint,
        mode="simple",
        standard_tags_enabled=True,
        taxonomy_hash="a" * 64,
        warn=lambda message: None,
    )
    mismatched = load_success_checkpoint(
        checkpoint,
        mode="simple",
        standard_tags_enabled=True,
        taxonomy_hash="b" * 64,
        warn=lambda message: None,
    )

    assert len(matching) == 1
    assert mismatched == {}


def test_simple_resume_loads_only_simple_checkpoint(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])

    run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=1,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([_simple_success_response]),
        printer=lambda message: None,
    )
    session = FakeSession()
    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=1,
            resume=True,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=lambda message: None,
    )

    assert summary.skipped_resume_count == 1
    assert session.calls == []


def test_batch_simple_mode_rejects_full_result(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=1,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([_success_response]),
        printer=lambda message: None,
    )

    failure = _read_jsonl(output / FAILED_ROWS_NAME)[0]
    assert summary.failed_count == 1
    assert failure["error_type"] == "mentor_result_validation_failed"
    assert _read_jsonl(output / SUCCESS_CHECKPOINT_RELATIVE) == []


def test_batch_full_mode_rejects_simple_result(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="full",
            batch_size=1,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([_simple_success_response]),
        printer=lambda message: None,
    )

    failure = _read_jsonl(output / FAILED_ROWS_NAME)[0]
    assert summary.failed_count == 1
    assert failure["error_type"] == "mentor_result_validation_failed"
    assert _read_jsonl(output / SUCCESS_CHECKPOINT_RELATIVE) == []


def test_full_resume_does_not_silently_load_simple_checkpoint(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1)])

    run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=1,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([_simple_success_response]),
        printer=lambda message: None,
    )
    printed: list[str] = []
    session = FakeSession([_success_response])

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="full",
            batch_size=1,
            resume=True,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=printed.append,
    )

    assert summary.skipped_resume_count == 0
    assert summary.success_count == 1
    assert len(session.calls) == 1
    assert any("mismatched full checkpoint" in message for message in printed)


def test_simple_batch_failure_falls_back_to_single_requests(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1), _mentor_row(2), _mentor_row(3)])
    session = FakeSession(
        [
            FakeResponse(payload={"choices": [{"message": {"content": "{}"}}]}),
            _simple_success_response,
            _simple_success_response,
            _simple_success_response,
        ]
    )
    printed: list[str] = []

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=3,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=printed.append,
    )

    assert summary.success_count == 3
    assert summary.failed_count == 0
    assert len(session.calls) == 4
    assert any("batch fallback size=3" in message for message in printed)
    assert len(_read_jsonl(output / SUCCESS_CHECKPOINT_RELATIVE)) == 3


def test_simple_batch_resume_skips_without_http_calls(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "outputs"
    _write_source(source, [_mentor_row(1), _mentor_row(2), _mentor_row(3)])

    run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=3,
            max_retries=0,
            no_excel=True,
        ),
        session=FakeSession([_simple_batch_success_response]),
        printer=lambda message: None,
    )
    session = FakeSession()

    summary = run_batch(
        BatchConfig(
            input_path=source,
            output_dir=output,
            mode="simple",
            batch_size=3,
            resume=True,
            max_retries=0,
            no_excel=True,
        ),
        session=session,
        printer=lambda message: None,
    )

    assert summary.skipped_resume_count == 3
    assert summary.success_count == 0
    assert summary.final_result_count == 3
    assert session.calls == []
